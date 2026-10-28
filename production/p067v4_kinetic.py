"""P067 revision 4: visible events drive the edit, not camera jitter.
Preserve the corrected Japanese voice, subtitle text, chronology and CC0 rock materials.
Every large scenic action is an artistic time-compressed reconstruction. No new facts.
"""
from pathlib import Path
import os,sys,json,math,shutil,zipfile,hashlib,subprocess
os.environ.setdefault('LIBGL_ALWAYS_SOFTWARE','1')
os.environ.setdefault('LP_NUM_THREADS','3')
os.environ.setdefault('OPENBLAS_NUM_THREADS','1')
import numpy as np
from PIL import Image,ImageDraw
import moderngl
import p067v3_phototexture as P
V=P.V;S=P.S;F=P.F
ROOT=Path('p067_v4').resolve()
for d in ['assets','render','review','source','deliverables','docs','audio']:(ROOT/d).mkdir(parents=True,exist_ok=True)
P.ROOT=V.ROOT=S.ROOT=F.ROOT=ROOT
F.PREFIX='heionreki_p067_v4'
BASECONFIG=P.config
BASECINEMA=P.PhotoCinema
BASETYPO=F.Typography
# A larger wave displacement has a matching analytic normal, so its silhouette moves.
S.VERT=S.VERT.replace('sea+.18*sin(a)+.12*sin(b)+.055*sin(c)+.026*sin(d)', 'sea+.32*sin(a)+.20*sin(b)+.085*sin(c)+.038*sin(d)')
S.VERT=S.VERT.replace('-.0432*cos(a)+.0168*cos(b)-.03355*cos(c)-.0247*cos(d)', '-.0768*cos(a)+.028*cos(b)-.05185*cos(c)-.0361*cos(d)')
S.VERT=S.VERT.replace('-.0216*cos(a)-.0372*cos(b)+.01815*cos(c)-.01898*cos(d)', '-.0384*cos(a)-.062*cos(b)+.02805*cos(c)-.02774*cos(d)')
S.FRAG=S.FRAG.replace('UV.y*.09-time*.032','UV.y*.09-time*.092')
S.FRAG=S.FRAG.replace('uniform sampler2D rocktex;', 'uniform float activepulse;uniform float landcool;uniform sampler2D rocktex;')
S.FRAG=S.FRAG.replace('col+=vec3(.42,.070,.009)*glow*heat*(.40+.60*ao);','col+=vec3(.54,.080,.009)*glow*heat*(.40+.60*ao)*(1.+activepulse*.70);')
S.FRAG=S.FRAG.replace('float crack=pow(1.-abs(2.*f.r-1.),28.)*smoothstep(.4,.75,g.b);','float crack=pow(1.-abs(2.*f.r-1.),28.)*smoothstep(.4,.75,g.b);crack*=mix(1.,smoothstep(landcool-.12,landcool+.12,g.r),.76);')
S.FRAG=S.FRAG.replace('else if(kind==10){',r'''else if(kind==12){
 vec4 f=texture(detailtex,L.xz*2.2+L.yy*.3);float crust=smoothstep(.39,.65,f.r);
 col=mix(tint.rgb,vec3(.025,.019,.015)*(.3+sunlit),crust*.62);
}else if(kind==13){
 col=tint.rgb*(.45+.65*max(0.,dot(n,l)));alpha=tint.a;
}else if(kind==10){''')
# Hot foreground matter stays locally bright; there is no full-frame exposure flash.

def pulse(t,at,width=.45):
 q=(t-at)/max(width,1e-5)
 return float(np.exp(-q*q*3.8))

def align_beam(a,b,r):
 a=np.asarray(a,float);b=np.asarray(b,float);direction=b-a;length=float(np.linalg.norm(direction))
 z=S.unit(direction);x=S.unit(np.cross([0,1,0],z))
 if np.linalg.norm(x)<.1:x=np.array([1.,0,0])
 y=np.cross(z,x);m=np.eye(4);m[:3,:3]=np.stack([x*r,y*r,z*max(length,.001)],axis=1);m[:3,3]=a;return m

_GEOM=None
def config(t,tl):
 global _GEOM
 c=BASECONFIG(t,tl);s=c['scene'];a,d=tl['boundaries'][s:s+2];sec=tl['cues'][s*2+1]['start'];u=c['u'];local=t-a
 c.update(scene_start=a,scene_end=d,second=sec,activepulse=0.,landcool=0.,kinetic_phase='')
 if s==0:
  c['activepulse']=max(pulse(t,a+.45,.40),pulse(t,a+1.35,.42))
  c['eruption']=.30+1.30*c['activepulse'];c['heat']=.70+.90*c['activepulse'];c['jetphase']=local
  if c['orbital']:c['landcool']=S.ramp(t,sec,d);c['heat']=1.25-.5*S.ramp(t,sec,d)
  c['kinetic_phase']='rim reveal, two ballistic surges, cut to whole forming planet'
 elif s==1:
  c['landcool']=.78*S.ramp(t,a,d);c['heat']=1.18-.50*S.ramp(t,a,d)
  c['activepulse']=.25+.35*np.sin(local*3.6)**2
  c['kinetic_phase']='cooling crust opens dark areas while the active surface turns; age reveal follows voice'
 elif s==2:
  travel_end=max(a+.4,sec-.20)
  q=float(np.clip((t-a-.12)/(travel_end-a-.12),0,1))
  # Acceleration is legible because the head sweeps a thickening illuminated path.
  c['time_progress']=q*q*(2.-q)
  c['arrival']=S.ramp(t,travel_end-.2,travel_end+.38)
  c['interval_reveal']=S.ramp(t,sec+.15,d-.12)
  c['activepulse']=pulse(t,travel_end,.32)
  c['kinetic_phase']='depart, illuminate milestones, ignite endpoint, pull back to show the interval'
 elif s==3:
  hit=S.keytime(tl,6,'インセキ',.8)+.36
  c['impact']=t-hit;c['impact_at']=hit
  c['activepulse']=max(pulse(t,sec+.37,.40),pulse(t,sec+1.38,.45))
  if t>=sec:c['eruption']=.35+1.5*c['activepulse'];c['heat']=.70+1.1*c['activepulse']
  c['jetphase']=t-sec
  c['kinetic_phase']='meteor approaches, impact ejecta spreads, lava surge, molten fragments cross the crater'
 elif s==4:
  c['activepulse']=max(pulse(t,a+.65,.55),pulse(t,a+1.72,.58))
  if not c['orbital']:c['eruption']=.50+.85*S.ramp(t,a,sec);c['jetphase']=local
  else:c['landcool']=.72;c['wet']=.25+.68*S.ramp(t,sec,d)
  c['kinetic_phase']='plume swells, rolls up the crater, then a growing atmospheric shell'
 elif s==5:
  # Three coherent advances rather than an almost static blue sphere.
  q=S.ramp(t,a+.05,d-.12)
  c['ocean']=.91*q;c['wet']=.55+.45*q;c['heat']=.60*(1-q);c['landcool']=.7
  c['kinetic_phase']='low basins flood, ocean joins, remaining high land becomes islands'
 elif s in [6,7,8]:
  c['wavephase']=local
  c['activepulse']=.35*pulse(t,sec+.45,.5);c['jetphase']=local
  if s==8 and t>=sec:
   q=S.ramp(t,sec,d)
   c['eye']=S.mix([-16,9.5,29],[-10,15.0,24],q).tolist();c['target']=[0,10,-18]
   c['event']='shore_break_then_reveal_volcanic_island'
  c['kinetic_phase']='foam front approaches grounded rock, visible spray rises, second crest passes; geography stays fixed'
 else:
  q=S.ramp(t,a+.1,d-.65)
  c['eye']=S.mix([-9,8.0,36],[-8,2.0,22],q).tolist();c['target']=S.mix([0,11,-18],[-7,2,5],q).tolist()
  c['wavephase']=local;c['activepulse']=0.;c['jetphase']=local
  c['kinetic_phase']='approach breaking surface, cross waterline, descend toward lit submerged relief'
 if not c['orbital']:
  if _GEOM is None:_GEOM=V.Geography()
  eye=np.asarray(c['eye'],float);floor=float(_GEOM.height(eye[0],eye[2]));eye[1]=max(eye[1],floor+.94);c['eye']=eye.tolist()
 return c

class Cinema(BASECINEMA):
 def __init__(self):
  super().__init__();self.mesh('kineticbeam',*V.tube_path([[0,0,0],[0,0,1]],1.,8))
  # Foamy crescent faces incoming water; it surrounds, rather than intersects, the rock.
  vv=[];ii=[]
  for k in range(65):
   ang=.12+(np.pi-.24)*k/64
   for rr in [2.85,3.38]:
    vv.append([rr*np.cos(ang),0,rr*np.sin(ang),0,1,0,k/64,(rr-2.85)/.53])
  for k in range(64):j=2*k;ii.extend([j,j+1,j+2,j+1,j+3,j+2])
  self.mesh('foamcrescent',np.array(vv,'f4'),np.array(ii,'i4'))
  self.shore_rocks=[]
  for x,z,sx,sz,top in [(-12.,24.,2.5,2.65,7.1),(-4.,8.,2.0,2.2,7.7)]:
   bottom=float(self.g.height(x,z))-.08;center=(bottom+top)/2.;sy=(top-bottom)/1.44
   m=S.matrix([x,center,z],[sx,sy,sz],.25);self.shore_rocks.append((m,[x,center,z],[sx*1.18,sy*.75,sz*1.18]))
  # Add the permanent shore rocks to the existing world-space shadow map.
  self.sf.use()
  for m,_,_ in self.shore_rocks:self.put(self.sp,'model',m);self.meshes['rock'][1].render()
  self.mesh('chronobeam',*V.tube_path([[0,0,0],[0,0,1]],1.,8))
 def setcommon(self,cfg,t,eye,vp):
  super().setcommon(cfg,t,eye,vp)
  self.put(self.prog,'activepulse',float(cfg.get('activepulse',0)))
  self.put(self.prog,'landcool',float(cfg.get('landcool',0)))
 def time_world(self,cfg,t):
  progress=float(cfg['time_progress']);a=V.A+[0,.85,0];b=V.B+[0,.85,0]
  self.draw('rail_left',11,tint=(.19,.43,.62,1));self.draw('rail_right',11,tint=(.48,.20,.045,1))
  for j in range(36):
   q=(j+.5)/36;p0=V.time_path(j/36);p1=V.time_path((j+1)/36)
   done=S.ramp(progress,q-.045,q+.012)
   if done>.001:self.draw('chronobeam',8,align_beam(p0,p1,.040+.025*done),tint=(1.6*done,.85*done,.16*done,1))
  for j in range(18):
   q=j/17;p=V.time_path(q);done=S.ramp(progress,q-.04,q+.01);sz=.8+.48*pulse(progress,q,.075)
   self.draw('tick',8,S.matrix(p,[sz,sz,sz]),tint=(.16+1.5*done,.27+.60*done,.38+.08*done,1))
  self.draw('planet',4,S.matrix(a,.235,angle=.30+t*.045))
  arrive=float(cfg.get('arrival',0));self.draw('sphere',10,S.matrix(b,.56+.11*arrive,angle=t*.08))
  self.draw('timehalo',8,S.matrix(a,[1.,1.,1.]),tint=(.70,.30,.07,1))
  self.draw('timehalo',8,S.matrix(b,[.98,.98,.98]),tint=(.10+.42*arrive,.23+.60*arrive,.35+.88*arrive,1))
  if progress<.995:
   for j in range(18):
    q=progress-j*.010
    if 0<=q<=1:self.draw('grain',8,S.matrix(V.time_path(q),.115*(1-j/23)),tint=(2.1,1.10,.21,1))
  # Following arrival, a broad travelling highlight measures the interval back to its origin.
  show=float(cfg.get('interval_reveal',0))
  if show>.03:
   q=1-show;pp=V.time_path(q)
   self.draw('tick',8,S.matrix(pp,[1.6,1.6,1.6]),tint=(.48,1.02,1.48,1))
 def jets(self,cfg,t):
  strength=float(cfg['eruption'])
  if strength<.08:return
  base=float(self.g.height(0,-18))+.55
  phase=float(cfg.get('jetphase',t))
  # Persistent ballistic histories: existing chunks do not teleport when power changes.
  for j in range(42):
   birth=math.floor((phase+j*.071)/1.55)*1.55-j*.071
   age=phase-birth
   if not 0<=age<=1.55:continue
   ang=j*2.39996;speed=1.6+.27*(j%7);vz=6.9+1.05*(j%5)
   p=np.array([speed*math.cos(ang)*age,base+vz*age-4.1*age*age,-18+speed*math.sin(ang)*age])
   if p[1]<float(self.g.height(p[0],p[2]))+.13:continue
   # The main readable surges use an increased number of luminous fragments.
   if j>18 and strength<.65:continue
   rad=.085+.022*(j%4)
   self.draw('rock',12,S.matrix(p,[rad,rad*1.5,rad],t*.8+j),tint=(3.8,.66,.035,1))
   vel=np.array([speed*math.cos(ang),vz-8.2*age,speed*math.sin(ang)])
   tail=p-S.unit(vel)*(.25+.24*min(strength,1.7))
   self.draw('kineticbeam',8,align_beam(tail,p,.018+.006*(j%3)),tint=(1.65,.20,.008,1))
 def impact(self,cfg,t):
  q=float(cfg.get('impact',-5));cx,cz=20.,10.;cy=float(self.g.height(cx,cz))
  if -.8<q<0:
   head=S.mix([31,37,29],[cx,cy+.18,cz],S.ease((q+.8)/.8));tail=head+[1.5,3.4,2.6]
   self.draw('rock',12,S.matrix(head,.42,angle=t),tint=(4.5,1.2,.10,1))
   self.draw('kineticbeam',8,align_beam(head,tail,.07),tint=(1.9,.30,.015,1))
  elif 0<=q<1.65:
   # Ejecta is localized; every fragment has a ballistic trajectory and ground test.
   for j in range(30):
    ang=j*2.399;speed=3.5+.55*(j%6);p=[cx+speed*q*np.cos(ang),cy+(4.+.65*(j%5))*q-4.5*q*q,cz+speed*q*np.sin(ang)]
    if p[1]<=float(self.g.height(p[0],p[2]))+.08:continue
    rr=.11+.035*(j%4);self.draw('rock',12,S.matrix(p,rr,angle=t+j),tint=(2.8,.50,.028,1))
 def surf(self,cfg,t):
  if not cfg['water_visible'] or cfg['orbital']:return
  sea=cfg['sea'];phase=float(cfg.get('wavephase',t));c=self.c
  c.enable(moderngl.BLEND);c.blend_func=(moderngl.SRC_ALPHA,moderngl.ONE_MINUS_SRC_ALPHA);c.depth_mask=False
  for k,(m,center,radii) in enumerate(self.shore_rocks):
   age=(phase-k*.48)%1.65;rise=math.sin(math.pi*min(age/1.1,1.));fade=max(0.,1-age/1.65)
   pos=np.array([center[0],sea+.09+.32*rise,center[2]])
   size=1+.18*age;self.draw('foamcrescent',13,S.matrix(pos,[size,1,size]),tint=(.82,.96,.96,.72*fade))
   if age<1.05:
    for j in range(15):
     ang=.12+(np.pi-.24)*j/14;r=2.9+.34*age
     p=[center[0]+r*np.cos(ang),sea+.15+(2.0+.08*(j%5))*age-2.5*age*age,center[2]+r*np.sin(ang)]
     self.draw('grain',13,S.matrix(p,[.035,.10,.035]),tint=(.80,.96,1.,.65*fade))
  c.depth_mask=True;c.disable(moderngl.BLEND)
 def world(self,cfg,t,with_water=True):
  if cfg.get('time_scene',False):self.time_world(cfg,t);return
  base=dict(cfg);base['eruption']=0.;base['impact']=-5.
  super().world(base,t,with_water)
  if not cfg['orbital']:
   for m,_,_ in self.shore_rocks:self.draw('rock',3,m)
   self.jets(cfg,t)
   if cfg['scene']==3 and t<cfg['second']:self.impact(cfg,t)
   if with_water:self.surf(cfg,t)
 def frame(self,cfg,t):
  im,qa=super().frame(cfg,t);eye=np.asarray(cfg['eye'])
  closest=[]
  if not cfg['orbital']:
   for _,center,radii in self.shore_rocks:
    gap=float(np.linalg.norm((eye-np.asarray(center))/np.asarray(radii)))
    closest.append(gap)
    assert gap>1.06,('Camera intersects added shore rock',t,gap)
  qa.update(kinetic_phase=cfg['kinetic_phase'],activepulse=float(cfg['activepulse']),landcool=float(cfg['landcool']),time_progress=float(cfg.get('time_progress',0)),ocean_fill=float(cfg.get('ocean',0)),shore_rock_camera_clearance=closest)
  return im,qa

class Typography(BASETYPO):
 def overlay(self,im,t,cfg,tl):
  if cfg['scene']!=2:return super().overlay(im,t,cfg,tl)
  # Keep the readable projected endpoint labels inherited from v3. Only the interval reveal changes.
  out=super().overlay(im,t,cfg,tl)
  return out

F.Cinema=Cinema;F.config=config;F.Typography=Typography

def review_package(kind='preflight'):
 # Private OOXML image contact sheets can be imported and rendered without publishing any media.
 from docx import Document
 from docx.shared import Inches,Pt
 dest=ROOT/'review_ooxml';dest.mkdir(exist_ok=True)
 doc=Document();sec=doc.sections[0];sec.top_margin=sec.bottom_margin=Inches(.3);sec.left_margin=sec.right_margin=Inches(.3)
 pics=sorted((ROOT/'review').glob('p067v2_s*.jpg'))
 if kind=='final':pics=sorted((ROOT/'review').glob('encoded_*.jpg'))
 for i,p in enumerate(pics):
  if i:doc.add_page_break()
  par=doc.add_paragraph(p.stem);par.paragraph_format.space_after=Pt(2)
  doc.add_picture(str(p),height=Inches(8.7))
 word=ROOT/'review/visual_review.docx';doc.save(word)
 with zipfile.ZipFile(word) as z:z.extractall(dest)
 print('PRIVATE_REVIEW_PAGES',len(pics),flush=True)

def assemble():
 F.assemble();tl=F.timeline();d=ROOT/'deliverables'
 qa=json.loads((ROOT/'review/verification.json').read_text())
 qa.update(revision='v4-kinetic',voice_and_caption_text_preserved=True,terrain_triangles=655360,added_grounded_shore_rocks=2,chronology='advancing illuminated rail, sequential milestone activation, endpoint ignition and interval reveal',large_action_policy='Volcanic surges, impact ejecta, evolving crust, ocean flooding and breaking shore waves. No constant camera shake.',visual_inspection_performed=False,continuous_preview_watched=False,auditory_listening_performed=False)
 F.dump(ROOT/'review/verification.json',qa)
 readings=json.loads((ROOT/'review/inherited_readings.json').read_text())
 assert all(x.get('query_reading_checked',False) for x in readings)
 notes='平穏暦全史・前史067 第4版\n八十億年前、別の世界が生まれた\n制作日：2026-10-29\n\n'
 notes+='今回の変更：火山弾の弾道と短い光跡、語りに連動する噴出の強弱、溶岩表面の速い流れと岩への照り返し、地面に接地する海岸岩と砕ける波、時系列の光路が順番に明るくなり太陽形成地点へ到達する演出。色・反射・凹凸の岩肌素材は第3版を保持。大型の形状・流れ・景観の変化を優先し、カメラの細かな振動や全画面の閃光を使わない。\n'
 notes+='高画質は1080×1920に直接描画。反射用の副描画は540×960。原音・字幕本文・発話時刻を維持し、前版の完成AAC音声を再圧縮せず合成。声と字幕の内容を勝手に変更しない。\n\n台本：\n'+'\n'.join(c['text'] for c in tl['cues'])
 notes+='\n\n設定と演出：\nオーウェルの約80億年前の形成、初期の衝突・火山活動、厚い二酸化炭素大気、世界規模の海が公式記事の軸。太陽形成の約46億年前との差は概算34億年。岩の位置、噴出の大きさ・間隔、海岸、波、時間を光路に置く構成は独自の再構成。年代軸は平穏暦に対応し、具体的景観は不穏暦の想像表現。火山・海岸・全球模型は縮尺を変えた説明表現。時間の光路は物理距離を示さず、二つの模型は異なる年代を比較する。二酸化炭素自体の色や、実際の火山周期を測った映像ではない。\n'
 notes+='\n出典：\nTSSL, Orwell, Geological Periods\nhttps://www.orionsarm.com/eg-article/546a05829e265\nTSSL, The Early Evolutionary History of Orwell\nhttps://www.orionsarm.com/eg-article/54749a15e55ef\nNASA Science, Sun Facts\nhttps://science.nasa.gov/sun/facts/\n科学・設定本文は第3版を継承。\n'
 notes+='\n素材：Poly Haven Rock Boulder Dry（CC0-1.0）、撮影Dimitrios Savva／処理Rico Cilliers。https://polyhaven.com/a/rock_boulder_dry ／ https://polyhaven.com/license 。原寸4K、描画材質2048角。火山岩向けに着色した景観表現であり原素材の岩種同定を意味しない。\n音声：VOICEVOX 四国めたん（ノーマル）。書体Noto Sans CJK JP。音楽・効果音は第3版を継承。フォント・声モデルは同梱しない。\n'
 notes+='\n確認済みの読み：\n'+json.dumps(readings,ensure_ascii=False,indent=2)
 notes+='\n\n自動検査と確認範囲：\n'+json.dumps(qa,ensure_ascii=False,indent=2)+'\n自動的な画素変化を見た目の迫力の保証と扱わない。実際に目視できた範囲は別の記録に残す。\n'
 (d/f'{F.PREFIX}_script_and_sources.txt').write_text(notes,encoding='utf-8')
 for name in ['p067v2_scene.py','p067v2_finish.py','p067v3_rocktime.py','p067v3_phototexture.py','p067v4_kinetic.py']:
  shutil.copy(Path(__file__).with_name(name),ROOT/'source'/name)
 (ROOT/'source/REEDIT.txt').write_text('p067_v4を作業ディレクトリへ展開。sourceの5本を同一フォルダに置く。numpy scipy Pillow moderngl==5.12.0 glcontext==3.0.0、ffmpegとEGL、日本語書体を用意。P067_FONTに書体パスを指定。python source/p067v4_kinetic.py shot N（0..9）。全場面の後assemble。original_previewには既存完成MP4を置き、AACをコピーする。raw voiceとmaster.wavはaudio内。素材はassets/polyhavenに同梱。',encoding='utf-8')
 review_package('final')
 archive=d/f'{F.PREFIX}_production.zip'
 if archive.exists():archive.unlink()
 manifest=[]
 for p in ROOT.rglob('*'):
  if p.is_file() and p.suffix.lower() not in ['.zip','.ttf','.otf','.ttc'] and 'review_ooxml' not in p.parts:
   manifest.append({'path':str(p.relative_to(ROOT)),'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()})
 F.dump(ROOT/'review/manifest.json',manifest)
 with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED,3) as z:
  for p in ROOT.rglob('*'):
   if p.is_file() and p!=archive and p.suffix.lower() not in ['.ttf','.otf','.ttc','.zip'] and 'review_ooxml' not in p.parts and 'voice_data' not in p.parts:z.write(p,str(p.relative_to(ROOT)))
 print('KINETIC_FINAL_QA',json.dumps(qa,ensure_ascii=False),flush=True)
 print('KINETIC_FILES',json.dumps([{'name':p.name,'bytes':p.stat().st_size} for p in d.iterdir()],ensure_ascii=False),flush=True)

if __name__=='__main__':
 mode=sys.argv[1]
 if mode=='prepare':P.acquire_assets();F.prepare()
 elif mode=='preview':
  for s in [0,2,3,5,6,8,9]:F.shot(s,True)
  review_package()
 elif mode=='testclip':F.shot(2);F.shot(3)
 elif mode=='shot':F.shot(int(sys.argv[2]))
 elif mode=='assemble':assemble()
 else:raise ValueError(mode)
