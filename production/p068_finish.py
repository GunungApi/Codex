"""Finish P068 from the actual generated voice and rendered 3-D frames."""
from pathlib import Path
import os,sys,json,math,subprocess,zipfile,hashlib,shutil
import numpy as np
from PIL import Image
from p068_scene import ROOT,W,H,FPS,Typography
PREFIX='heionreki_p068_v1'
def dump(p,x):Path(p).write_text(json.dumps(x,ensure_ascii=False,indent=2),encoding='utf-8')
def run(args,**kw):return subprocess.run(args,check=True,**kw)
def ts(t,ass=False):
 if ass:return f'{int(t//3600)}:{int(t//60)%60:02}:{int(t)%60:02}.{int(t%1*100):02}'
 n=round(t*1000);return f'{n//3600000:02}:{n//60000%60:02}:{n//1000%60:02},{n%1000:03}'
def subtitles(tl):
 typ=Typography();ja=[];en=[];ass=['[Script Info]','ScriptType: v4.00+','PlayResX: 1080','PlayResY: 1920','WrapStyle: 2','ScaledBorderAndShadow: yes','[V4+ Styles]','Format: Name, Fontname, Fontsize, PrimaryColour, SecondaryColour, OutlineColour, BackColour, Bold, Italic, Underline, StrikeOut, ScaleX, ScaleY, Spacing, Angle, BorderStyle, Outline, Shadow, Alignment, MarginL, MarginR, MarginV, Encoding','Style: Default,Noto Sans CJK JP,64,&H00F8F7F2,&H00EED775,&HBB180E03,&HBB180E03,1,0,0,0,100,100,0,0,3,13,0,5,60,60,200,1','[Events]','Format: Layer, Start, End, Style, Name, MarginL, MarginR, MarginV, Effect, Text']
 for c in tl['cues']:
  a=max(0,c['start']-.06);b=min(tl['duration'],c['end']+.14);lines=typ.lines(c['text']);j='\n'.join(lines)
  ja.append(f"{c['id']}\n{ts(a)} --> {ts(b)}\n{j}\n");en.append(f"{c['id']}\n{ts(a)} --> {ts(b)}\n{c['en']}\n")
  tx=r'{\move(540,1568,540,1552,0,180)\fad(140,120)}'+r'\N'.join(lines);ass.append(f'Dialogue: 0,{ts(a,True)},{ts(b,True)},Default,,0,0,0,,{tx}')
 for suf,ls in [('ja.srt',ja),('en.srt',en),('ja.ass',ass)]:(ROOT/'deliverables'/f'{PREFIX}_{suf}').write_text('\n'.join(ls),encoding='utf-8')

def review(mode='preflight'):
 from docx import Document
 from docx.shared import Inches,Pt
 doc=Document();sec=doc.sections[0];sec.page_width=Inches(6.15);sec.page_height=Inches(11.25);sec.top_margin=sec.bottom_margin=Inches(.22);sec.left_margin=sec.right_margin=Inches(.25)
 if mode=='preflight':
  pics=[]
  for s in range(10):
   found=sorted((ROOT/'review').glob(f'p068_s{s:02}_f*.jpg'))
   if found:pics.extend([found[min(1,len(found)-1)],found[min(3,len(found)-1)]])
 else:pics=sorted((ROOT/'review').glob('encoded_*.jpg'))
 for i,p in enumerate(pics):
  if i:doc.add_page_break()
  para=doc.add_paragraph(p.stem);para.paragraph_format.space_after=Pt(3);para.runs[0].font.size=Pt(9)
  doc.add_picture(str(p),width=Inches(5.6))
 dest=ROOT/'review'/f'{mode}.docx';doc.save(dest)
 folder=ROOT/'review'/f'{mode}_ooxml'
 if folder.exists():shutil.rmtree(folder)
 folder.mkdir()
 with zipfile.ZipFile(dest) as z:z.extractall(folder)
 print('PRIVATE_VISUAL_REVIEW',mode,len(pics),dest.stat().st_size,flush=True)

def finish():
 tl=json.loads((ROOT/'timeline.json').read_text());N=round(tl['duration']*FPS);subtitles(tl)
 stats=[]
 for s in range(10):stats+=json.loads((ROOT/'review'/f'states_{s:02}.json').read_text())
 assert [s['frame'] for s in stats]==list(range(N));assert not any(s['issues'] for s in stats)
 paths=[ROOT/'render'/f'shot_{s:02}.mp4' for s in range(10)];assert all(p.exists() for p in paths)
 concat=ROOT/'render/concat.txt';concat.write_text(''.join("file '"+str(p)+"'\n" for p in paths))
 picture=ROOT/'render/full_picture.mp4';run(['ffmpeg','-y','-v','error','-f','concat','-safe','0','-i',str(concat),'-c','copy',str(picture)])
 high=ROOT/'deliverables'/f'{PREFIX}_high_1080x1920.mp4';light=ROOT/'deliverables'/f'{PREFIX}_light_720x1280.mp4'
 run(['ffmpeg','-y','-v','error','-i',str(picture),'-i',str(ROOT/'audio/master.wav'),'-map','0:v:0','-map','1:a:0','-c:v','copy','-c:a','aac','-b:a','192k','-ar','48000','-t',str(tl['duration']),'-movflags','+faststart',str(high)])
 run(['ffmpeg','-y','-v','error','-i',str(high),'-vf','scale=720:1280:flags=lanczos','-c:v','libx264','-preset','fast','-crf','23','-c:a','copy','-threads','3','-movflags','+faststart',str(light)])
 probes=[]
 for p in [high,light]:
  probe=json.loads(subprocess.check_output(['ffprobe','-v','error','-show_streams','-show_format','-of','json',str(p)]));v=next(x for x in probe['streams'] if x['codec_type']=='video');a=next(x for x in probe['streams'] if x['codec_type']=='audio');assert int(v['nb_frames'])==N;assert v['r_frame_rate']=='24/1';assert a['sample_rate']=='48000';run(['ffmpeg','-v','error','-i',str(p),'-map','0:v','-map','0:a','-f','null','-']);probes.append(probe)
 vals=np.array([s['delta'] for s in stats]);window=np.convolve(vals,np.ones(48)/48,'valid')
 measurement=subprocess.run(['ffmpeg','-hide_banner','-i',str(high),'-af','loudnorm=I=-16:TP=-1.5:LRA=9:print_format=json','-f','null','-'],capture_output=True,text=True,check=True).stderr
 loud=json.loads(measurement[measurement.rfind('{'):measurement.rfind('}')+1])
 # Verify the encoded audio is identical in both delivery sizes.
 audiohash=[]
 for p in [high,light]:audiohash.append(subprocess.check_output(['ffmpeg','-v','error','-i',str(p),'-map','0:a:0','-c','copy','-f','hash','-hash','sha256','-']).decode().strip())
 assert audiohash[0]==audiohash[1]
 for i,c in enumerate(tl['cues']):
  at=(c['start']+c['end'])/2;out=ROOT/'review'/f'encoded_{i:02}.jpg';run(['ffmpeg','-y','-v','error','-ss',str(at),'-i',str(high),'-frames:v','1','-q:v','2',str(out)])
 report={'episode':'P068','title':'地球より先に、この海では生命が生きていた','duration':tl['duration'],'fps':FPS,'frames':N,'native_resolution':[W,H],'caption_count':len(tl['cues']),'max_query_speech_gap':tl['max_query_speech_gap'],'loudness':loud,'audio_identical_between_delivery_sizes':True,'minimum_camera_clearance':min(s['camera_clearance'] for s in stats),'main_cell_geometry_issues':[],'minimum_2_second_frame_change':float(window.min()),'weak_motion_windows_below_0_015':np.flatnonzero(window<.015).tolist(),'visual_inspection_performed':False,'continuous_preview_watched':False,'auditory_listening_performed':False,'scope':'World-space bounding volumes of principal cells and mineral floor; projected centers, all rendered frames. Nutrient absorption intentionally crosses the modeled shell. Pixel change is only a motion aid, not proof of perceptual quality. Cell morphology, colors, marine micro-landscape, magnification and light pulses are original explanatory reconstructions, not official organism anatomy or bioluminescence.'}
 dump(ROOT/'review/verification.json',report);dump(ROOT/'review/probes.json',probes)
 notes=['平穏暦全史・前史068 第1版','地球より先に、この海では生命が生きていた','','構成：P067の海から顕微鏡的な縮尺へ。単純な微生物、周囲の有機物の利用、酸素を使わない暮らし、暖かな生息環境、初期生命の時代の入口を見せる。','範囲：約73億年前からという時代区分の入口。正確な起源の日時や特定の起源の場所を確定しない。後代の光合成、核相当構造、動物、森林は含めない。','','ナレーション：VOICEVOX 四国めたん（ノーマル）。原音20区間、実際の音声の長さと合成器の発音時刻から字幕と場面を同期。漢字の読み辞書と照合記録を同梱。','映像：1080×1920で直接描画した独自の3D地形・鉱物・微生物模型。膜の質感、奥行き、光の透過感、被写界深度、材料の受け渡しを設計。','細胞の具体的な姿、色、内部の粒、海底の景観は説明用の創作。縮尺を場面ごとに切替。光の走行は生きる働きを示す記号であり、発光生物という設定の追加ではない。','体内の粒は材料の記号。核や葉緑体、特定の地球生物の器官として描いていない。','音楽・効果音：今回のための合成音。外部の楽曲、画像、音声モデル、フォントファイルは制作ZIPに含めない。','','台本：']
 notes += [f"{c['id']:02}  {c['start']:.3f}–{c['end']:.3f}  {c['text']}" for c in tl['cues']]
 notes+=['','出典（2026-10-30確認）：','TSSL／Orion’s Arm Universe Project — Orwell, Geological Periods','https://www.orionsarm.com/eg-article/546a05829e265','TSSL／Orion’s Arm Universe Project — The Early Evolutionary History of Orwell','https://www.orionsarm.com/eg-article/54749a15e55ef','制作計画：平穏暦全史 第3巻 詳細制作計画 第1.0版 P068。','本作はユウ・星野ルナによる解説・再現映像。公式の原文やイラストを再掲載していない。','','検査結果：',json.dumps(report,ensure_ascii=False,indent=2)]
 (ROOT/'deliverables'/f'{PREFIX}_script_and_sources.txt').write_text('\n'.join(notes),encoding='utf-8')
 for name in ['p068_scene.py','p068_refine.py','p068_finish.py']:
  shutil.copy(Path(__file__).with_name(name),ROOT/'source'/name)
 (ROOT/'source/REEDIT.txt').write_text('p068_v1として展開。sourceを同じディレクトリに置き、P068_FONTでNoto Sans CJK JPの場所を指定。python source/p068_scene.py shot 0〜9、python source/p068_finish.py assemble。音声は同梱の原音とmaster.wavを保持する。再生成には公式VOICEVOXエンジンが必要。フォント・エンジン・声モデルは同梱しない。\n',encoding='utf-8')
 review('final')
 archive=ROOT/'deliverables'/f'{PREFIX}_production.zip'
 with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED,3) as z:
  for p in ROOT.rglob('*'):
   if p.is_file() and p!=archive and p.suffix.lower() not in ['.ttf','.ttc','.otf','.docx','.pyc'] and not any(x.endswith('_ooxml') for x in p.parts) and 'voice_data' not in p.parts:
    z.write(p,str(p.relative_to(ROOT)))
 print('P068_FINAL',json.dumps(report,ensure_ascii=False),flush=True)
 print('DELIVERABLES',json.dumps([{'name':p.name,'bytes':p.stat().st_size} for p in (ROOT/'deliverables').iterdir()],ensure_ascii=False),flush=True)
if __name__=='__main__':
 if sys.argv[1]=='assemble':finish()
 elif sys.argv[1]=='review':review(sys.argv[2] if len(sys.argv)>2 else 'preflight')
 else:raise ValueError(sys.argv[1])
