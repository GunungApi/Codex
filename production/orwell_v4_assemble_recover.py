"""Assembly-only repair: normalize ffmpeg JPEGs for python-docx image signatures.
Finished 1080p shots are recovered unchanged; no scene is redrawn by this step.
"""
from pathlib import Path
import sys,io,zipfile,shutil,json
from PIL import Image
from docx import Document
from docx.shared import Inches,Pt
import orwell_v4_finalpass as R
K=R.K;F=R.F;ROOT=R.ROOT

def review_package(kind='final'):
 dest=ROOT/'review_ooxml'
 if dest.exists():shutil.rmtree(dest)
 dest.mkdir()
 doc=Document();sec=doc.sections[0];sec.top_margin=sec.bottom_margin=Inches(.30);sec.left_margin=sec.right_margin=Inches(.30)
 pics=sorted((ROOT/'review').glob('encoded_[0-9][0-9].jpg')) if kind=='final' else sorted((ROOT/'review').glob('p067v2_s*.jpg'))
 if kind=='final':assert len(pics)==20,len(pics)
 for i,p in enumerate(pics):
  if i:doc.add_page_break()
  para=doc.add_paragraph(p.stem);para.paragraph_format.space_after=Pt(2)
  # Re-encode to a well-defined PNG signature; do not alter finished MP4 pixels.
  buf=io.BytesIO();im=Image.open(p).convert('RGB');im.save(buf,format='PNG');buf.seek(0)
  doc.add_picture(buf,height=Inches(8.7))
 target=ROOT/'review/final_encoded_visual_review.docx';doc.save(target)
 with zipfile.ZipFile(target) as z:z.extractall(dest)
 print('NORMALIZED_REVIEW_IMAGES',len(pics),flush=True)

K.review_package=review_package
if __name__=='__main__':
 if sys.argv[1]=='prepare':R.P.acquire_assets();F.prepare()
 elif sys.argv[1]=='assemble':
  R.assemble()
  shutil.copy(__file__,ROOT/'source'/Path(__file__).name)
  print('RECOVERED_ASSEMBLY_COMPLETE',flush=True)
 else:raise ValueError(sys.argv[1])
